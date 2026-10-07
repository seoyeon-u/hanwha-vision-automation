"""Offline behavioral tests. No real API requests and no user data changes."""
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

import httpx
from openai import RateLimitError, BadRequestError
from openpyxl import load_workbook
from PIL import Image

import config as cfg
import main


def fixture(image_path="m1-1-001.png"):
    return {
        "imagePath": image_path, "valid": False,
        "shapes": [{"label": "전차", "points": [[782, 319], [1170, 475]],
                    "시선방향": "측면", "무장여부": True,
                    "무기타입": ["포/포탑_1", "원격무장_1"], "무기방향": ["이외", "이외"]}],
    }


def response(text="One tank is facing to the side. Its mounted weapons have unspecified directions."):
    return SimpleNamespace(status="completed", output_text=text)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        settings = {
            "IMAGE_DIR": self.root / "images", "JSON_DIR": self.root / "json",
            "OUTPUT_FILE": self.root / "output" / "descriptions.xlsx",
            "CHECKPOINT_FILE": self.root / "output" / "checkpoint.jsonl",
            "ENV_FILE": self.root / ".env", "RESUME": True,
            "MAX_ATTEMPTS": 3, "SAVE_EVERY": 1,
        }
        self.patcher = patch.multiple(cfg, **settings)
        self.patcher.start()
        self.addCleanup(self.patcher.stop)
        for directory in [cfg.IMAGE_DIR, cfg.JSON_DIR, cfg.OUTPUT_FILE.parent]:
            directory.mkdir(parents=True)
        cfg.ENV_FILE.write_text("OPENAI_API_KEY=test-key\n", encoding="utf-8")

    def write_json(self, name="sample.json", value=None):
        path = cfg.JSON_DIR / name
        path.write_text(json.dumps(fixture() if value is None else value, ensure_ascii=False), encoding="utf-8")
        return path

    def test_real_schema_and_false_valid(self):
        original, annotation = main.parse_annotation(self.write_json())
        self.assertEqual(original, "m1-1-001.png")
        self.assertEqual(annotation["class_counts"], {"tank": 1})
        self.assertIs(annotation["metadata"]["valid"], False)
        tank = annotation["objects"][0]
        self.assertEqual(tank["body_direction"], "side")
        self.assertEqual(tank["weapon_types"], ["gun/turret_1", "remote weapon system_1"])
        self.assertEqual(tank["weapon_directions"], ["other", "other"])
        self.assertEqual(tank["bbox"], [782, 319, 1170, 475])

    def test_counts_and_individual_person_attributes(self):
        value = fixture()
        value["shapes"] = [
            {"label": "군인", "points": [[1, 2], [3, 4]], "시선방향": "정면", "무장여부": True, "무기타입": ["소총"]},
            {"label": "군인", "points": [[5, 6], [7, 8]], "시선방향": "측면", "무장여부": False},
        ]
        _, annotation = main.parse_annotation(self.write_json(value=value))
        self.assertEqual(annotation["class_counts"], {"soldier": 2})
        self.assertEqual(annotation["objects"][0]["weapon_types"], ["rifle"])
        self.assertIs(annotation["objects"][1]["armed"], False)

    def test_relative_basename_absolute_and_ambiguity(self):
        nested = cfg.IMAGE_DIR / "nested"
        nested.mkdir()
        image = nested / "sample.png"
        image.write_bytes(b"matching only")
        matcher = main.ImageMatcher(cfg.IMAGE_DIR)
        self.assertEqual(matcher.resolve("nested\\sample.png"), image)
        self.assertEqual(matcher.resolve("old/location/sample.png"), image)
        self.assertEqual(matcher.resolve(str(image)), image)
        (cfg.IMAGE_DIR / "sample.png").write_bytes(b"second matching file")
        matcher = main.ImageMatcher(cfg.IMAGE_DIR)
        with self.assertRaises(ValueError):
            matcher.resolve("old/sample.png")
        with self.assertRaises(FileNotFoundError):
            matcher.resolve("missing.png")

    def test_image_conversion_and_validation(self):
        bmp = cfg.IMAGE_DIR / "sample.bmp"
        Image.new("RGB", (3, 3)).save(bmp)
        self.assertTrue(main.image_data_url(bmp).startswith("data:image/png;base64,"))
        bmp.write_bytes(b"broken image")
        with self.assertRaises(OSError):
            main.image_data_url(bmp)

    @patch("main.time.sleep")
    def test_empty_response_retry_and_combined_input(self, sleep):
        client = MagicMock()
        client.responses.create.side_effect = [response(" "), response()]
        _, annotation = main.parse_annotation(self.write_json())
        description = main.generate_description(client, annotation, "data:image/png;base64,AAAA")
        self.assertTrue(description.startswith("One tank"))
        self.assertEqual(client.responses.create.call_count, 2)
        request = client.responses.create.call_args.kwargs
        self.assertEqual([item["type"] for item in request["input"][0]["content"]], ["input_text", "input_image"])
        self.assertIn('"valid":false', request["input"][0]["content"][0]["text"])
        self.assertFalse(request["store"])
        sleep.assert_called_once()

    @patch("main.time.sleep")
    def test_rate_limit_retry_and_permanent_error(self, sleep):
        client = MagicMock()
        request = httpx.Request("POST", "https://api.openai.com/v1/responses")
        limit = RateLimitError("test", response=httpx.Response(429, request=request), body=None)
        client.responses.create.side_effect = [limit, response()]
        main.generate_description(client, {}, "data:image/png;base64,AAAA")
        self.assertEqual(client.responses.create.call_count, 2)
        client.reset_mock()
        client.responses.create.side_effect = BadRequestError("test", response=httpx.Response(400, request=request), body=None)
        with self.assertRaises(BadRequestError):
            main.generate_description(client, {}, "data:image/png;base64,AAAA")
        self.assertEqual(client.responses.create.call_count, 1)

    @patch("main.time.sleep")
    def test_exhausted_empty_response_and_format_validation(self, sleep):
        client = MagicMock()
        client.responses.create.return_value = response("")
        with self.assertRaises(main.InvalidDescription):
            main.generate_description(client, {}, "data:image/png;base64,AAAA")
        self.assertEqual(client.responses.create.call_count, cfg.MAX_ATTEMPTS)
        for text in ["Description: A tank.", "1. A tank.", "```A tank.```", "전차 한 대."]:
            with self.subTest(text=text), self.assertRaises(main.InvalidDescription):
                main.validate_description(text)

    @patch("main.time.sleep")
    def test_quota_errors_are_not_retried(self, sleep):
        request = httpx.Request("POST", "https://api.openai.com/v1/responses")
        for code in main.QUOTA_CODES:
            with self.subTest(code=code):
                client = MagicMock()
                error = RateLimitError("test", response=httpx.Response(429, request=request),
                                       body={"code": code, "type": "insufficient_quota"})
                client.responses.create.side_effect = error
                with self.assertRaises(RateLimitError):
                    main.generate_description(client, {}, "data:image/png;base64,AAAA")
                self.assertEqual(client.responses.create.call_count, 1)
                self.assertEqual(main.api_error_details(error), (code, "insufficient_quota"))
        sleep.assert_not_called()

    def test_retry_after_minimum_is_not_capped(self):
        request = httpx.Request("POST", "https://api.openai.com/v1/responses")
        error = RateLimitError("test", response=httpx.Response(429, request=request, headers={"Retry-After": "90"}),
                              body={"code": "rate_limit_exceeded", "type": "rate_limit_error"})
        self.assertEqual(main.retry_delay(error, 1), 90.0)
        error.response.headers["Retry-After"] = "invalid"
        self.assertGreater(main.retry_delay(error, 1), 0)

    def test_quota_error_blocks_later_requests_but_exports_all_rows(self):
        self.write_json("a.json", fixture("first.png"))
        self.write_json("b.json", fixture("second.png"))
        for name in ["first.png", "second.png"]:
            Image.new("RGB", (3, 3)).save(cfg.IMAGE_DIR / name)
        request = httpx.Request("POST", "https://api.openai.com/v1/responses")
        client = MagicMock()
        client.responses.create.side_effect = RateLimitError(
            "test", response=httpx.Response(429, request=request),
            body={"error": {"code": "insufficient_quota", "type": "insufficient_quota"}},
        )
        with patch("main.OpenAI", return_value=client), patch("main.time.sleep") as sleep:
            self.assertEqual(main.run(), 1)
            sleep.assert_not_called()
        self.assertEqual(client.responses.create.call_count, 1)
        workbook = load_workbook(cfg.OUTPUT_FILE)
        try:
            rows = list(workbook["description"].values)
            self.assertEqual(rows[1:], [("first.png", None), ("second.png", None)])
        finally:
            workbook.close()

    def test_excel_original_paths_and_checkpoint_recovery(self):
        original = "=folder\\sample.png"
        self.assertTrue(main.save_excel([(original, "One tank is visible."), ("failed.png", "")]))
        workbook = load_workbook(cfg.OUTPUT_FILE)
        try:
            sheet = workbook["description"]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), ("imagePath", "description"))
            self.assertEqual(sheet["A2"].value, original)
            self.assertEqual(sheet["A2"].data_type, "s")
        finally:
            workbook.close()
        with cfg.CHECKPOINT_FILE.open("w", encoding="utf-8") as handle:
            handle.write('{"interrupted":')
        main.append_checkpoint("recovered.png", "One soldier is standing.")
        results = main.load_results()
        self.assertEqual(results[original], "One tank is visible.")
        self.assertEqual(results["recovered.png"], "One soldier is standing.")
        self.assertNotIn("failed.png", results)

    def test_pipeline_continues_then_resumes_without_api_calls(self):
        self.write_json("a_missing.json", fixture("missing.png"))
        self.write_json("b_good.json")
        self.write_json("c_duplicate.json")
        (cfg.JSON_DIR / "d_corrupt.json").write_text("broken", encoding="utf-8")
        invalid = fixture("invalid.png")
        invalid["shapes"] = "invalid"
        self.write_json("e_invalid.json", invalid)
        Image.new("RGB", (3, 3)).save(cfg.IMAGE_DIR / "m1-1-001.png")
        client = MagicMock()
        client.responses.create.return_value = response()
        with patch("main.OpenAI", return_value=client):
            self.assertEqual(main.run(), 1)
        self.assertEqual(client.responses.create.call_count, 1)
        workbook = load_workbook(cfg.OUTPUT_FILE)
        try:
            rows = list(workbook["description"].values)
            self.assertEqual(len(rows), 5)  # Header + 4 readable imagePaths.
            self.assertEqual(rows[2][1], rows[3][1])
            self.assertIsNone(rows[1][1])
            self.assertEqual(rows[4][0], "invalid.png")
        finally:
            workbook.close()
        with patch("main.OpenAI") as constructor:
            self.assertEqual(main.run(), 1)
            constructor.assert_not_called()

    def test_fresh_runs_call_api_for_every_annotation(self):
        self.write_json("a.json")
        self.write_json("b.json")
        Image.new("RGB", (3, 3)).save(cfg.IMAGE_DIR / "m1-1-001.png")
        main.save_excel([("m1-1-001.png", "Old saved description.")])
        main.append_checkpoint("m1-1-001.png", "Old checkpoint description.")
        client = MagicMock()
        descriptions = [f"One tank is visible in scene {number}." for number in range(4)]
        client.responses.create.side_effect = [response(text) for text in descriptions]
        with patch.object(cfg, "RESUME", False), patch("main.OpenAI", return_value=client):
            self.assertEqual(main.run(), 0)
            self.assertEqual(main.run(), 0)
        self.assertEqual(client.responses.create.call_count, 4)
        workbook = load_workbook(cfg.OUTPUT_FILE)
        try:
            rows = list(workbook[cfg.SHEET_NAME].values)
            self.assertEqual([row[1] for row in rows[1:]], descriptions[2:])
        finally:
            workbook.close()

    def test_locked_excel_keeps_success_checkpoint(self):
        main.append_checkpoint("good.png", "One tank is visible.")
        with patch("main.os.replace", side_effect=PermissionError("locked")):
            self.assertFalse(main.save_excel([("good.png", "One tank is visible.")]))
        self.assertEqual(main.load_results()["good.png"], "One tank is visible.")


if __name__ == "__main__":
    unittest.main()
