"""Generate one English caption per annotation JSON with resumable Excel export."""
from __future__ import annotations

import base64
import math
from collections import Counter, defaultdict
from datetime import datetime
from email.utils import parsedate_to_datetime
import io
import json
import logging
import os
from pathlib import Path, PureWindowsPath
import random
import re
import sys
import time

from dotenv import dotenv_values
from openai import APIConnectionError, APIStatusError, APITimeoutError, OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from PIL import Image

import config as cfg
from prompt import SYSTEM_PROMPT, build_annotation_text

LOGGER = logging.getLogger("descriptions")
STRUCTURAL_KEYS = {"label", "points", "group_id", "shape_type", "mask"}
VEHICLES = {"light tactical vehicle", "tank", "self-propelled artillery", "armored vehicle", "civilian vehicle"}
QUOTA_CODES = {
    "insufficient_quota", "credit_balance_exhausted", "billing_hard_limit_reached",
    "billing_not_active", "organization_spend_limit_exceeded",
    "project_spend_limit_exceeded", "organization_usage_limit_exceeded",
}


def api_error_details(exc: APIStatusError) -> tuple[str, str]:
    """Extract machine-readable fields only; never log bodies, keys or headers."""
    body = getattr(exc, "body", None)
    if isinstance(body, dict) and isinstance(body.get("error"), dict):
        body = body["error"]
    fields = []
    for name in ("code", "type"):
        value = body.get(name) if isinstance(body, dict) else None
        value = value or getattr(exc, name, None)
        fields.append(value if isinstance(value, str) and re.fullmatch(r"[a-z_]{1,80}", value) else "unknown")
    return tuple(fields)


def is_quota_error(exc: APIStatusError) -> bool:
    code, error_type = api_error_details(exc)
    return code in QUOTA_CODES or error_type == "insufficient_quota"


def api_error_hint(exc: APIStatusError) -> str:
    if is_quota_error(exc):
        return "API 크레딧/사용 한도 오류입니다. 해당 API 프로젝트의 Billing과 Limits를 확인한 뒤 다시 실행하세요."
    if exc.status_code == 429:
        return "API 요청/토큰 제한입니다. 잠시 후 재실행하거나 프로젝트의 모델별 Limits를 확인하세요."
    if exc.status_code == 401:
        return "프로젝트 .env의 OPENAI_API_KEY가 유효한지 확인하세요."
    if exc.status_code == 403:
        return "API 프로젝트 및 모델 접근 권한을 확인하세요."
    return "API 오류 코드와 요청 설정을 확인하세요."


def retry_delay(exc: Exception, attempt: int) -> float:
    delay = min(cfg.RETRY_MAX_SECONDS, cfg.RETRY_BASE_SECONDS * 2 ** (attempt - 1) + random.uniform(0, 1))
    if isinstance(exc, APIStatusError):
        raw = exc.response.headers.get("retry-after")
        if raw:
            try:
                server_delay = float(raw)
            except ValueError:
                try:
                    server_delay = parsedate_to_datetime(raw).timestamp() - time.time()
                except (ValueError, TypeError, OverflowError):
                    server_delay = 0.0
            if math.isfinite(server_delay):
                # The server's minimum wait must not be shortened by our backoff cap.
                delay = max(delay, server_delay)
    return delay


def wait_before_retry(delay: float) -> None:
    while delay > 0:
        chunk = min(delay, 30.0)
        time.sleep(chunk)
        delay -= chunk


def setup_logging() -> None:
    cfg.LOG_DIR.mkdir(parents=True, exist_ok=True)
    LOGGER.setLevel(logging.INFO)
    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    for handler in [logging.StreamHandler(sys.stdout), logging.FileHandler(
        cfg.LOG_DIR / f"run_{datetime.now():%Y%m%d_%H%M%S}.log", encoding="utf-8"
    )]:
        handler.setFormatter(formatter)
        LOGGER.addHandler(handler)


def map_values(value, mapping: dict):
    if isinstance(value, list):
        return [map_values(item, mapping) for item in value]
    return mapping.get(value, value) if isinstance(value, str) else value


def translate_weapon(value):
    if isinstance(value, list):
        return [translate_weapon(item) for item in value]
    if isinstance(value, str):
        name, separator, suffix = value.rpartition("_")
        if separator and suffix.isdigit():
            return cfg.WEAPON_NAMES.get(name, name) + separator + suffix
        return cfg.WEAPON_NAMES.get(value, value)
    return value


def parse_annotation(path: Path) -> tuple[str, dict]:
    with path.open(encoding="utf-8-sig") as handle:
        raw = json.load(handle)
    if not isinstance(raw, dict):
        raise ValueError("JSON root must be an object")
    image_path = raw.get("imagePath")
    if not isinstance(image_path, str) or not image_path.strip():
        raise ValueError("imagePath must be a nonempty string")
    shapes = raw.get("shapes")
    if not isinstance(shapes, list):
        raise ValueError("Expected LabelMe shapes array")
    objects = []
    for number, shape in enumerate(shapes, 1):
        if not isinstance(shape, dict) or not isinstance(shape.get("label"), str) or not shape["label"].strip():
            raise ValueError(f"shapes[{number - 1}] must have a nonempty label")
        points = shape.get("points")
        if not isinstance(points, list) or not points:
            raise ValueError(f"shapes[{number - 1}] is missing points/bbox")
        for point in points:
            if not isinstance(point, list) or len(point) != 2 or any(
                isinstance(v, bool) or not isinstance(v, (int, float)) for v in point
            ):
                raise ValueError(f"shapes[{number - 1}] has malformed points")
        cls = cfg.CLASS_NAMES.get(shape["label"], shape["label"])
        attrs = {k: v for k, v in shape.items() if k not in STRUCTURAL_KEYS}
        # Keep original attributes (including flags/custom fields) as well as English conveniences.
        item = {"object_id": number, "class": cls, "original_class": shape["label"],
                "bbox": [min(p[0] for p in points), min(p[1] for p in points),
                         max(p[0] for p in points), max(p[1] for p in points)],
                "attributes": attrs}
        if "시선방향" in shape:
            item["body_direction" if cls in VEHICLES else "gaze_direction"] = map_values(shape["시선방향"], cfg.DIRECTIONS)
        for source, target in [("본체방향", "body_direction"), ("무기방향", "weapon_directions")]:
            if source in shape:
                item[target] = map_values(shape[source], cfg.DIRECTIONS)
        if "무장여부" in shape:
            item["armed"] = shape["무장여부"]
        if "무기타입" in shape:
            item["weapon_types"] = translate_weapon(shape["무기타입"])
        types, directions = item.get("weapon_types"), item.get("weapon_directions")
        if isinstance(types, list) and isinstance(directions, list) and len(types) != len(directions):
            raise ValueError(f"shapes[{number - 1}] weapon type/direction lengths differ")
        objects.append(item)
    # Exclude embedded image bytes; the actual image is sent separately. Preserve all other metadata.
    metadata = {k: v for k, v in raw.items() if k not in {"shapes", "imageData"}}
    return image_path, {"metadata": metadata, "class_counts": dict(Counter(o["class"] for o in objects)), "objects": objects}


class ImageMatcher:
    def __init__(self, root: Path):
        self.root = root.resolve()
        self.by_name = defaultdict(list)
        for path in sorted(self.root.rglob("*")):
            if path.is_file() and path.suffix.lower() in cfg.IMAGE_EXTENSIONS:
                self.by_name[path.name.casefold()].append(path)

    def resolve(self, original: str) -> Path:
        normalized = original.replace("\\", "/")
        path = Path(normalized)
        windows_path = PureWindowsPath(original)
        absolute = path.is_absolute() or windows_path.is_absolute()
        candidate = path if absolute else self.root / path
        if candidate.is_file():
            return candidate
        matches = self.by_name.get(PureWindowsPath(normalized).name.casefold(), [])
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            raise ValueError(f"Ambiguous image basename for {original!r}: {len(matches)} matches; use correct relative path")
        raise FileNotFoundError(f"Image not found: {original!r}")


def image_data_url(path: Path) -> str:
    # Pillow verifies the input and converts unsupported BMP/TIFF to API-compatible PNG.
    with Image.open(path) as image:
        image.load()
        if getattr(image, "n_frames", 1) > 1:
            raise ValueError("Animated images are not supported; provide a still image")
        mime = {"JPEG": "image/jpeg", "PNG": "image/png", "WEBP": "image/webp", "GIF": "image/gif"}.get(image.format)
        if mime:
            content = path.read_bytes()
        else:
            buffer = io.BytesIO()
            image.convert("RGB").save(buffer, format="PNG")
            content, mime = buffer.getvalue(), "image/png"
    return f"data:{mime};base64," + base64.b64encode(content).decode("ascii")


class InvalidDescription(ValueError):
    pass


def validate_description(text: str) -> str:
    text = text.strip()
    if not text:
        raise InvalidDescription("Empty API response")
    if len(text) > 32767:
        raise InvalidDescription("Description exceeds Excel cell limit")
    if re.search(r"[\uac00-\ud7a3]", text) or not re.search(r"[A-Za-z]", text):
        raise InvalidDescription("Description is not English")
    if re.match(r"(?i)^(description\s*:|\d+[.)]\s|[-*#>]\s|[\[{\"])", text) or "```" in text:
        raise InvalidDescription("Response contains a heading, numbering, JSON or Markdown")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text):
        raise InvalidDescription("Response contains invalid Excel control characters")
    return " ".join(text.split())


def generate_description(client: OpenAI, annotation: dict, image_url: str) -> str:
    for attempt in range(1, cfg.MAX_ATTEMPTS + 1):
        try:
            response = client.responses.create(
                model=cfg.MODEL, instructions=SYSTEM_PROMPT, store=False,
                max_output_tokens=cfg.MAX_OUTPUT_TOKENS,
                input=[{"role": "user", "content": [
                    {"type": "input_text", "text": build_annotation_text(annotation)},
                    {"type": "input_image", "image_url": image_url, "detail": cfg.IMAGE_DETAIL},
                ]}],
            )
            if response.status != "completed":
                raise InvalidDescription(f"API response status: {response.status}")
            return validate_description(response.output_text)
        except (APIConnectionError, APITimeoutError, APIStatusError, InvalidDescription) as exc:
            retryable = not isinstance(exc, APIStatusError) or exc.status_code in {408, 409, 429} or exc.status_code >= 500
            if isinstance(exc, APIStatusError) and is_quota_error(exc):
                retryable = False
            if not retryable or attempt == cfg.MAX_ATTEMPTS:
                raise
            delay = retry_delay(exc, attempt)
            # Do not log API response bodies, headers or keys.
            code, error_type = api_error_details(exc) if isinstance(exc, APIStatusError) else ("n/a", "n/a")
            LOGGER.warning("Retry %d/%d in %.1fs (%s, code=%s, type=%s)", attempt, cfg.MAX_ATTEMPTS, delay, type(exc).__name__, code, error_type)
            wait_before_retry(delay)
    raise RuntimeError("Retry loop exhausted")


def load_results() -> dict[str, str]:
    results = {}
    if not cfg.RESUME:
        return results
    if cfg.OUTPUT_FILE.exists():
        workbook = load_workbook(cfg.OUTPUT_FILE, read_only=True, data_only=False)
        try:
            sheet = workbook[cfg.SHEET_NAME]
            rows = sheet.iter_rows(values_only=True)
            if tuple(next(rows, ())) != ("imagePath", "description"):
                raise ValueError("Existing workbook has incorrect headers")
            for row in rows:
                if len(row) >= 2 and isinstance(row[0], str) and isinstance(row[1], str) and row[1].strip():
                    results[row[0]] = row[1]
        finally:
            workbook.close()
    if cfg.CHECKPOINT_FILE.exists():
        with cfg.CHECKPOINT_FILE.open(encoding="utf-8") as handle:
            for number, line in enumerate(handle, 1):
                if not line.strip():
                    continue
                try:
                    item = json.loads(line)
                    if not isinstance(item.get("imagePath"), str) or not isinstance(item.get("description"), str) or not item["description"].strip():
                        raise ValueError("Invalid checkpoint entry")
                    results[item["imagePath"]] = item["description"]
                except (ValueError, TypeError, AttributeError):
                    LOGGER.warning("Ignoring damaged checkpoint line %d", number)
    return results


def append_checkpoint(image_path: str, description: str) -> None:
    with cfg.CHECKPOINT_FILE.open("a", encoding="utf-8") as handle:
        # Leading newline separates a possibly interrupted final line from the new entry.
        handle.write("\n" + json.dumps({"imagePath": image_path, "description": description}, ensure_ascii=False) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def save_excel(rows: list[tuple[str, str]]) -> bool:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = cfg.SHEET_NAME
    sheet.append(["imagePath", "description"])
    for row_number, (image_path, description) in enumerate(rows, 2):
        sheet.append([image_path, description])
        # Store paths exactly as text, including paths starting with '='.
        for cell in (sheet.cell(row_number, 1), sheet.cell(row_number, 2)):
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 45
    sheet.column_dimensions["B"].width = 110
    temp = cfg.OUTPUT_FILE.with_name(cfg.OUTPUT_FILE.stem + ".tmp.xlsx")
    try:
        workbook.save(temp)
        os.replace(temp, cfg.OUTPUT_FILE)
        return True
    except OSError as exc:
        LOGGER.error("Excel save failed (%s). Close Excel and rerun; successful results remain in checkpoint.", type(exc).__name__)
        return False
    finally:
        workbook.close()


def run() -> int:
    for directory in (cfg.IMAGE_DIR, cfg.JSON_DIR, cfg.OUTPUT_FILE.parent, cfg.CHECKPOINT_FILE.parent):
        directory.mkdir(parents=True, exist_ok=True)
    if cfg.SAVE_EVERY < 1 or cfg.MAX_ATTEMPTS < 1 or cfg.SAVE_INTERVAL_SECONDS <= 0:
        raise ValueError("SAVE_EVERY/MAX_ATTEMPTS must be positive; SAVE_INTERVAL_SECONDS must be > 0")
    files = sorted(p for p in cfg.JSON_DIR.rglob("*") if p.is_file() and p.suffix.lower() == ".json")
    if not files:
        LOGGER.info("No annotation JSON files. Put files in %s", cfg.JSON_DIR)
        return 0
    results = load_results()
    if not cfg.RESUME and cfg.CHECKPOINT_FILE.exists():
        # Start a new checkpoint without deleting old successful results.
        backup = cfg.CHECKPOINT_FILE.with_name(f"checkpoint_{time.time_ns()}.bak.jsonl")
        cfg.CHECKPOINT_FILE.replace(backup)
    matcher = ImageMatcher(cfg.IMAGE_DIR)
    rows = []
    success = failed = skipped = processed = 0
    client = None
    blocked_api_error = None
    last_save = time.monotonic()
    interrupted = False
    try:
        for index, json_path in enumerate(files, 1):
            LOGGER.info("[%d/%d %.1f%%] %s", index, len(files), 100 * index / len(files), json_path.relative_to(cfg.JSON_DIR))
            row_index = None
            try:
                image_path, annotation = parse_annotation(json_path)
                if len(image_path) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", image_path):
                    raise ValueError("imagePath cannot be represented unchanged in an Excel cell")
                row_index = len(rows)
                rows.append((image_path, ""))
                if cfg.RESUME and image_path in results:
                    rows[row_index] = (image_path, results[image_path])
                    skipped += 1
                    LOGGER.info("Resume: already generated %s", image_path)
                else:
                    matched = matcher.resolve(image_path)
                    if blocked_api_error is not None:
                        raise blocked_api_error
                    if client is None:
                        # Explicitly read this one .env value; do not fall back to process environment.
                        key = dotenv_values(cfg.ENV_FILE, interpolate=False).get("OPENAI_API_KEY")
                        if not key or key.strip() in {"your_api_key", "..."}:
                            raise ValueError("Set OPENAI_API_KEY in the project .env file")
                        client = OpenAI(api_key=key.strip(), max_retries=0, timeout=cfg.REQUEST_TIMEOUT)
                    description = generate_description(client, annotation, image_data_url(matched))
                    append_checkpoint(image_path, description)
                    results[image_path] = description
                    rows[row_index] = (image_path, description)
                    success += 1
                    LOGGER.info("API 호출 성공: %s — description 생성 및 체크포인트 저장 완료", image_path)
            except Exception as exc:
                failed += 1
                if row_index is None:
                    try:
                        with json_path.open(encoding="utf-8-sig") as handle:
                            source = json.load(handle)
                        readable_path = source.get("imagePath") if isinstance(source, dict) else None
                        if isinstance(readable_path, str) and readable_path.strip() and len(readable_path) <= 32767 and not re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", readable_path):
                            rows.append((readable_path, ""))
                    except (OSError, ValueError):
                        pass
                if isinstance(exc, APIStatusError):
                    code, error_type = api_error_details(exc)
                    LOGGER.error("Failed %s: %s (HTTP %s, code=%s, type=%s). %s", json_path.name, type(exc).__name__, exc.status_code, code, error_type, api_error_hint(exc))
                    if is_quota_error(exc) or exc.status_code in {401, 403}:
                        blocked_api_error = exc
                else:
                    LOGGER.error("Failed %s: %s: %s", json_path.name, type(exc).__name__, exc)
            processed = index
            if processed % cfg.SAVE_EVERY == 0 or time.monotonic() - last_save >= cfg.SAVE_INTERVAL_SECONDS:
                save_excel(rows)
                last_save = time.monotonic()
    except KeyboardInterrupt:
        interrupted = True
        LOGGER.warning("Interrupted. Saving processed rows; rerun to resume.")
    finally:
        # Preserve recovered successes beyond the portion visited before interruption.
        if interrupted:
            existing_paths = {image_path for image_path, _ in rows}
            rows.extend((p, d) for p, d in results.items() if p not in existing_paths)
        saved = save_excel(rows)
        if client is not None:
            client.close()
        LOGGER.info("Completed: success=%d, failed=%d, skipped=%d, processed=%d/%d", success, failed, skipped, processed, len(files))
        LOGGER.info("Excel: %s", cfg.OUTPUT_FILE)
    return 130 if interrupted else (1 if failed or not saved else 0)


def main() -> int:
    setup_logging()
    try:
        return run()
    except Exception as exc:
        LOGGER.error("Startup failed: %s: %s", type(exc).__name__, exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
